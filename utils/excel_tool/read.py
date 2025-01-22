import copy
import re
import openpyxl
from decimal import Decimal
import openpyxl.worksheet.worksheet
import re
from openpyxl.utils import column_index_from_string
from util.excel_tool.excel_formula_tool import handle_formula

excel_pos_re = re.compile("\$?(?P<col>[A-Z]+)\$?(?P<row>\d+)")


def number_index_add(number):
    while True:
        number += yield number


class AssociatedField:
    def __init__(self, header_direction="row", header_position=(1, 1), rule=None, no_key_continue=True, define_key_rule=True):
        """
        用来定义读取的规则，ExcelReadTool类中就是通过这个类定义的规则来读取对应的字段的
        :param header_direction:
            这个参数表示是以每行或者是每列的方式来进行便利读取
        :param header_position:
            这个参数表示参数的第一个位置,其中第一个参数是表示的是行，第二个参数表示的是列的位置
        :param rule:

        :param no_key_continue:
        :param define_key_rule:
        """
        if header_direction not in ['col', 'row']:
            raise Exception("head_direction must is 'col' or 'row'")
        if rule is None:
            rule = {}
        if header_position[0] < 1 or header_position[1] < 1:
            raise Exception("header x or y must greater than 0")
        self._rule = rule
        self._header_direction = header_direction
        self._header_position = header_position
        self._no_key_continue = no_key_continue
        self._primary_key = None
        self._keys = []
        self._value_key = []
        self._datas = []
        self._define_key_rule = define_key_rule
        self._user_key_rule = None

    def get_primary_key(self):
        return self._primary_key

    def get_continue_keys(self):
        continue_keys = copy.deepcopy(self._keys)
        if self._primary_key is not None:
            continue_keys.append(self._primary_key)
        return continue_keys

    def get_no_key_continue(self):
        return self._no_key_continue

    def get_rule(self):
        return self._rule

    def get_header_point(self):
        return self._header_position

    def get_header_direction(self):
        return self._header_direction

    def add_field(self, field_name, last_step=1, is_primary_key=False, key=False):
        if self._rule.get(field_name) is not None:
            raise Exception("field is repeat")
        self._rule[field_name] = {
            "step": last_step
        }
        if is_primary_key:
            if self._primary_key is None:
                self._primary_key = field_name
            else:
                raise Exception("primary key only one")
        elif key:
            self._keys.append(field_name)
        else:
            self._value_key.append(field_name)


class ExcelReadTool:
    def __init__(self, excel_path):
        self._wb = openpyxl.load_workbook(excel_path, data_only=False)
        self._excel_path = excel_path
        self._property = {}

    def get_all_sheet_name(self):
        return self._wb.sheetnames

    def read(self, sheet_name, ass_field: AssociatedField):
        if sheet_name not in self.get_all_sheet_name():
            raise Exception("not have this sheet {},{}".format(sheet_name, self.get_all_sheet_name()))
        work_sheet = self._wb[sheet_name]
        sheet_property = self._property.get(sheet_name)
        if sheet_property is None:
            sheet_property = {}
            self._property[sheet_name] = sheet_property

        merged_list = sheet_property.get("merged_list")
        if merged_list is None:
            merged_list = self.get_excel_merged_list(work_sheet)
            sheet_property["merged_list"] = merged_list
        start_point = ass_field.get_header_point()
        if ass_field.get_header_direction() == "row":
            end = work_sheet.max_row + 1
            start = start_point[0]
        else:
            end = work_sheet.max_column + 1
            start = start_point[1]
        datas = []
        for data_index in range(start, end):
            data = {}
            if ass_field.get_header_direction() == "row":
                row = number_index_add(data_index)
                row.send(None)
                col = number_index_add(start_point[1])
                col.send(None)
                step = col
            else:
                row = number_index_add(start_point[0])
                row.send(None)
                col = number_index_add(data_index)
                col.send(None)
                step = row

            for rule_key, rule_property in ass_field.get_rule().items():
                next_step = rule_property.get("step")
                step.send(next_step)
                data_value = self.read_cell_value(work_sheet, row.send(0), col.send(0), merged_list)
                data[rule_key] = data_value

            if ass_field.get_no_key_continue():
                continue_keys = ass_field.get_continue_keys()
                need_continue = False
                for continue_key in continue_keys:
                    if data.get(continue_key) is None:
                        need_continue = True
                        break
                if need_continue:
                    continue
            datas.append(data)
        return datas

    @classmethod
    def create_sheet(cls, wb, sheet_index, sheet_name, all_datas=None, parameter_variations=None):
        sheet = wb.create_sheet(index=sheet_index, title=sheet_name)
        row = 1
        col = 1
        cls._sheet_wirte_datas(sheet, all_datas, parameter_variations, row, col)

    @classmethod
    def _sheet_wirte_datas(cls, sheet, values, parameter_variations, row=1, index=1):
        if isinstance(parameter_variations, dict):
            primary_key = parameter_variations.get("primary key", None)
            if primary_key is None:
                for variations_k, variations_v in parameter_variations.items():
                    sheet.cell(row=row, column=index).value = variations_k
                    values_v = values.get(variations_k, None)
                    if values_v:
                        row = cls._sheet_wirte_datas(sheet, values.get(variations_k), variations_v, row, index + 1)
                    # row += 1
            else:
                for variations_k, variations_str in parameter_variations.items():
                    if variations_k == "primary key":
                        continue
                    if not isinstance(variations_str, str):
                        continue
                    try:
                        variation_obj = eval(variations_str)
                    except Exception as e:
                        # print(variations_str)
                        raise Exception(e)
                    if variation_obj.index == {}:
                        continue
                    sheet.cell(row=row, column=index + variation_obj.index).value = variations_k
                row += 1
                for value in values:
                    if not isinstance(value, dict):
                        raise "存储的数据的字典，必须放在列表下"
                    for value_k, value_v in value.items():
                        variations_str = parameter_variations.get(value_k, None)
                        if variations_str is None:
                            continue
                        variation_obj = eval(variations_str)
                        sheet.cell(row=row, column=index + variation_obj.index).value = value_v
                    row += 1
                else:
                    row += 1
        return row

    @staticmethod
    def _excel_merged_split_letter_number(str_data):
        for index, s in enumerate(str_data):
            if s.isnumeric():
                return ord(str_data[:index]) - 64, int(str_data[index:])

    @classmethod
    def get_excel_merged_list(cls, worksheet):
        merged_list = []
        for merged_cell in worksheet.merged_cells:
            merged_cell = merged_cell.__str__()
            merged_start, merged_end = merged_cell.split(":", 1)
            merged_start_col, merged_start_row = cls._excel_merged_split_letter_number(merged_start)
            merged_end_col, merged_end_row = cls._excel_merged_split_letter_number(merged_end)
            merged = [merged_start_row, merged_start_col, merged_end_row, merged_end_col]
            merged_list.append(merged)
        merged_list = sorted(merged_list, key=lambda x: (x[0], x[1]))
        return merged_list

    @staticmethod
    def __handle_cell_data(data):
        if isinstance(data, str):
            return data.strip()
        return data

    @classmethod
    def read_cell_value(cls, worksheet, row, col, merged_list):
        cell_obj = worksheet.cell(row=row, column=col)
        if len(dir(cell_obj)) == 47:
            merged_start = cls.search_merged([row, col], 0, len(merged_list) - 1, merged_list)
            if merged_start is None:
                result_value = worksheet.cell(row=row, column=col).value
                return cls.__handle_cell_data(result_value)
            else:
                result_value = worksheet.cell(row=merged_start[0], column=merged_start[1]).value
                return cls.__handle_cell_data(result_value)
        else:
            if cell_obj.data_type == "f":
                result_value = worksheet.cell(row=row, column=col).value
                result_value = result_value.lstrip("=")
                transform_data = ""
                while True:
                    excel_pos_re_result = excel_pos_re.search(result_value)
                    if excel_pos_re_result is None:
                        transform_data = "{}{}".format(transform_data, result_value)
                        break
                    excel_cell_row = excel_pos_re_result.group("row")
                    excel_cell_col = excel_pos_re_result.group("col")
                    excel_cell_row_number = int(excel_cell_row)
                    excel_cell_col_number = int(column_index_from_string(excel_cell_col))
                    cell_value = cls.read_cell_value(worksheet, excel_cell_row_number, excel_cell_col_number, merged_list)
                    transform_data = "{}{}{}".format(transform_data, result_value[:excel_pos_re_result.start()], cell_value)
                    result_value = result_value[excel_pos_re_result.end():]

                result_data = handle_formula(transform_data)
                return result_data
            else:
                result_value = worksheet.cell(row=row, column=col).value
            return cls.__handle_cell_data(result_value)

    @classmethod
    def row_col_is_in_merged_list(cls, row, col, merged_data: str):

        merged_start, merged_end = merged_data.split(":", 1)

        merged_start_col, merged_start_row = cls._excel_merged_split_letter_number(merged_start)
        merged_end_col, merged_end_row = cls._excel_merged_split_letter_number(merged_end)
        if merged_start_row <= row <= merged_end_row and merged_start_col <= col <= merged_end_col:
            return merged_start_row, merged_start_col
        return None

    @staticmethod
    def __compare_merged(position, merged):
        if position[0] > merged[2]:
            return True
        elif position[0] < merged[0]:
            return False
        else:
            if position[1] > merged[3]:
                return True
            elif position[1] < merged[1]:
                return False
            else:
                return tuple(merged[:2])

    @classmethod
    def search_merged(cls, position, start_index, end_index, mergeds_list):
        mid_index = (end_index + start_index) // 2
        search_data = cls.__compare_merged(position, mergeds_list[mid_index])
        if end_index - start_index <= 1:
            if isinstance(search_data, bool) and mid_index == end_index:
                return None
            elif isinstance(search_data, bool):
                search_data = cls.__compare_merged(position, mergeds_list[end_index])
                if isinstance(search_data, bool):
                    return None
            return search_data
        if search_data is True:
            return cls.search_merged(position, mid_index, end_index, mergeds_list)
        elif search_data is False:
            return cls.search_merged(position, start_index, mid_index, mergeds_list)
        else:
            return search_data

    def change_cell_vale(self, work_sheet, row, col, value):
        work_sheet.cell(row=row, column=col).value = value
        return None

    def change_save(self, excel_path=None):
        if excel_path is None:
            excel_path = self._excel_path

        self._wb.save(excel_path)
